"""Recalculate earlier blink/whisker Wilson and nose-angle cluster intervals."""
import argparse
import json
import math
import random
import re
from collections import Counter, defaultdict
from pathlib import Path


def wilson(k, n):
    z = 1.959963984540054
    p = k / n
    d = 1 + z * z / n
    center = (p + z * z / (2 * n)) / d
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return [center - half, center + half]


def pct(vals, fraction):
    vals = sorted(vals)
    position = (len(vals) - 1) * fraction
    lo, hi = math.floor(position), math.ceil(position)
    return vals[lo] + (position - lo) * (vals[hi] - vals[lo])


def video_agreement(path, task):
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    groups = defaultdict(list)
    for row in data:
        groups[re.sub(r"_\d+$", "", row["video_id"])].append(row)
    details = []
    for video, rows in groups.items():
        labels = {int(row["label"]) for row in rows}
        if len(labels) != 1:
            raise ValueError(f"Mixed labels: {video}")
        label = labels.pop()
        if task == "blink" and label == 0:
            continue
        counts = Counter(int(row["pred_top1"]) for row in rows)
        winners = [pred for pred, count in counts.items() if count == max(counts.values())]
        exact = {int(pred == label) for pred in winners}
        within = {int(abs(pred - label) <= 1) for pred in winners}
        details.append({"video_id": video, "clips": len(rows), "reference": label, "majority_candidates": winners,
                        "exact": exact.pop() if len(exact) == 1 else None,
                        "within_one": within.pop() if len(within) == 1 else None})
    summary = {}
    for metric in ("exact", "within_one"):
        values = [row[metric] for row in details]
        if None in values:
            summary[metric] = {"status": "omitted: tie rule affects this metric"}
        else:
            k, n = sum(values), len(values)
            summary[metric] = {"successes": k, "n": n, "estimate": k / n, "wilson_95_ci": wilson(k, n)}
    return {"source": path.name, "video_details": details, "wilson_independent_video_assumption": summary}


def nose(path, resamples, seed):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["after_sample"]
        if sheet["B3"].value != "Ground Truth" or sheet["K3"].value != "Prediction":
            raise ValueError("Unexpected workbook layout")
        groups = defaultdict(list)
        pairs = []
        for row in range(4, 11):
            animal = "old animal" if row == 4 else f"rat {(row - 5) % 3 + 1}"
            # Rows 5-7 are trial 1, rows 8-10 are trial 2.
            for offset in range(6):
                manual_cell = sheet.cell(row, 4 + offset)
                auto_cell = sheet.cell(row, 13 + offset)
                manual, auto = manual_cell.value, auto_cell.value
                if not isinstance(manual, (int, float)) or not isinstance(auto, (int, float)):
                    continue
                error = abs(auto - manual)
                detail = {"animal": animal, "manual_cell": manual_cell.coordinate, "automated_cell": auto_cell.coordinate,
                          "manual": manual, "automated": auto, "absolute_error": error, "within_5_strict": int(error < 5)}
                groups[animal].append(detail)
                pairs.append(detail)
    finally:
        workbook.close()
    if len(pairs) != 17 or len(groups) != 4:
        raise ValueError(f"Expected 17 paired measurements from four animals; found {len(pairs)} from {len(groups)}")
    k = sum(p["within_5_strict"] for p in pairs)
    rng = random.Random(seed)
    animals = sorted(groups)
    mae_draws, rate_draws = [], []
    for _ in range(resamples):
        sampled = [pair for _ in animals for pair in groups[animals[rng.randrange(len(animals))]]]
        mae_draws.append(sum(p["absolute_error"] for p in sampled) / len(sampled))
        rate_draws.append(sum(p["within_5_strict"] for p in sampled) / len(sampled))
    return {"source": path.name, "complete_pairs": len(pairs), "animals": {animal: len(rows) for animal, rows in groups.items()},
            "mae_degrees": sum(p["absolute_error"] for p in pairs) / len(pairs),
            "mae_animal_cluster_bootstrap_95_ci": [pct(mae_draws, .025), pct(mae_draws, .975)],
            "within_5_strict": {"successes": k, "n": len(pairs), "estimate": k / len(pairs),
                                "wilson_independent_pair_95_ci": wilson(k, len(pairs)),
                                "animal_cluster_bootstrap_95_ci": [pct(rate_draws, .025), pct(rate_draws, .975)]},
            "pairs": pairs, "note": "Only four animals; cluster intervals are exploratory. Wilson assumes independent pairs."}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=Path(__file__).parent / "data" / "previous_ci")
    parser.add_argument("--output", type=Path, default=Path(__file__).parent / "results" / "other_ci.json")
    parser.add_argument("--resamples", type=int, default=2000)
    parser.add_argument("--seed", type=int, default=20261001)
    args = parser.parse_args()
    if args.resamples < 1:
        parser.error("resamples must be positive")
    result = {"resamples": args.resamples, "seed": args.seed,
              "blink": video_agreement(args.source_dir / "results_blink.json", "blink"),
              "whisker": video_agreement(args.source_dir / "results_whisker_backup.json", "whisker"),
              "nose": nose(args.source_dir / "nose_degree.xlsx", args.resamples, args.seed)}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Saved {args.output}")
    for task in ("blink", "whisker"):
        print(task, result[task]["wilson_independent_video_assumption"])
    print("nose", {key: value for key, value in result["nose"].items() if key != "pairs"})


if __name__ == "__main__":
    main()
